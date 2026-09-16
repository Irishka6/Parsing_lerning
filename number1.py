import csv
import requests
from bs4 import BeautifulSoup
from datetime import datetime
from collections import defaultdict
import matplotlib.pyplot as plt
from docx import Document
from docx.shared import Inches

monthly_count = defaultdict(int)
all_news_data = []  # <-- ДОБАВЛЕНО: для хранения всех новостей (для docx)

url = 'https://media.kpfu.ru/news?kn%5B0%5D=Международное%20сотрудничество&created='
response = requests.get(url)
html = response.text
soup = BeautifulSoup(html, 'html.parser')

with open('mai_news.csv', 'w', newline='', encoding='utf-8-sig') as file:
    writer = csv.writer(file)
    writer.writerow(['Ссылка', 'Заголовок', 'Дата'])

    page = 1
    while True:
        if page == 1:
            current_url = 'https://media.kpfu.ru/news?kn%5B0%5D=Международное%20сотрудничество&created='

        if page > 1:
            try:
                response = requests.get(current_url, timeout=10)
            except requests.exceptions.RequestException as e:
                print(f"Ошибка запроса (страница {page}): {e}")
                print("Прерываем парсинг, переходим к отчёту...")
                break
            html = response.text
            soup = BeautifulSoup(html, 'html.parser')
        news = soup.select('.newsItem')

        if not news:
            break

        for new in news:
            # --- ДОБАВЛЕНО: защита от None ---
            link_tag = new.select_one('.boldLink')
            date_tag_elem = new.select_one('.newsItem-date')
            if not link_tag or not date_tag_elem:
                continue
            # ---------------------------------

            htt = 'https://media.kpfu.ru/news?kn%5B0%5D=Международное%20сотрудничество&created=' + link_tag.get('href')
            title_tag = link_tag.text.strip()
            date_tag = date_tag_elem.text.strip()
            print(htt, title_tag, date_tag)
            writer.writerow([htt, title_tag, date_tag])

            # --- ДОБАВЛЕНО: сохраняем в список для docx ---
            all_news_data.append({'link': htt, 'title': title_tag, 'date': date_tag})
            # -----------------------------------------------

            try:
                date_obj = datetime.strptime(date_tag, '%d.%m.%Y')
                month_key = date_obj.strftime('%B %Y')
                monthly_count[month_key] += 1
            except:
                pass

        # --- ДОБАВЛЕНО: защита от None у кнопки "next" ---
        next_btn = soup.select_one('.pager__item.pager__item--next a')
        last_btn = soup.select_one('.pager__item.pager__item--last a')
        if not next_btn or not last_btn:
            break
        # -------------------------------------------------

        current_url = 'https://media.kpfu.ru/news?kn%5B0%5D=Международное%20сотрудничество&created=' + next_btn.get('href')
        print(current_url)
        if current_url == soup.select_one('.pager__item.pager__item--last a').get('href'):
            break
        page += 1

print("\n" + "=" * 50)
print("СТАТИСТИКА ПО МЕСЯЦАМ")
print("=" * 50)
print(f"{'Месяц':<20} {'Количество новостей':<20}")
print("-" * 50)
for month, count in sorted(monthly_count.items(), key=lambda x: datetime.strptime(x[0], '%B %Y')):
    print(f"{month:<20} {count:<20}")
print("-" * 50)
print(f"{'ИТОГО:':<20} {sum(monthly_count.values()):<20}")
print("=" * 50)

# ============================================================
# ДОБАВЛЕНО: построение графика распределения по месяцам
# ============================================================
if monthly_count:
    sorted_months = sorted(monthly_count.keys(), key=lambda x: datetime.strptime(x, '%B %Y'))
    counts = [monthly_count[m] for m in sorted_months]

    plt.figure(figsize=(10, 6))
    plt.bar(sorted_months, counts, color='skyblue')
    plt.title('Распределение новостей по месяцам (Международное сотрудничество)')
    plt.xlabel('Месяц')
    plt.ylabel('Количество новостей')
    plt.xticks(rotation=45, ha='right')
    plt.tight_layout()
    plt.savefig('news_chart.png')
    print("\nГрафик сохранён в news_chart.png")

    # ============================================================
    # ДОБАВЛЕНО: генерация docx-отчёта с графиком и таблицей
    # ============================================================
    doc = Document()
    doc.add_heading('Отчёт по новостям (Международное сотрудничество)', 0)

    doc.add_heading('График распределения по месяцам', level=1)
    doc.add_picture('news_chart.png', width=Inches(6.0))

    doc.add_heading('Таблица публикаций', level=1)
    table = doc.add_table(rows=1, cols=3)
    table.style = 'Table Grid'
    hdr = table.rows[0].cells
    hdr[0].text = 'Ссылка'
    hdr[1].text = 'Заголовок'
    hdr[2].text = 'Дата'

    for item in all_news_data:
        row = table.add_row().cells
        row[0].text = item['link']
        row[1].text = item['title']
        row[2].text = item['date']

    doc.save('news_report.docx')
    print("DOCX-отчёт сохранён в news_report.docx")